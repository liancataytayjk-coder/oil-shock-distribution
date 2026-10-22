"""Build slide decks in the Lian signature style with python-pptx.

Every slide is drawn on the template's Blank layout with fixed geometry, so
every deck lines up the same way: title, a one-sentence key message, a thin
rule, the body, and a footer with the short title and the slide number.

    import sys; sys.path.insert(0, "<skill>/scripts")
    from lian_deck import Deck

    d = Deck(short_title="Critical review: Agbola and Acupan (2010)")
    d.title_slide("Critical Review: ...", ["Agbola, F. W. & Acupan, A. B. (2010)",
                  "*Economic Systems*, 34, 386–396"], byline="Prepared by: **Lian Cataytay**")
    s = d.slide("1. Research problem and gap", "Emigration became permanent, but its causes were not well established.")
    body, bottom = d.split(d.body(), 1.0)          # body area and a strip for a callout
    left, right = d.columns(2, body)
    d.text(s, left, ["In the 1970s ...", "It became a **permanent** part ..."], heading="The problem")
    d.text(s, right, ["The causes ..."], heading="The gap")
    d.callout(s, bottom, "Research question:", "How do economic, demographic and political factors affect emigration?")
    d.takeaway("Push factors at home shape Filipino emigration.", ["**Unemployment raises emigration** ..."])
    d.closing()
    d.save("deck.pptx")

Inline markup in any text: **bold**, *italic*, ^{superscript}, _{subscript}.
A list item may be a string or (level, string) with level 0 or 1.
Text that would overflow its box is shrunk 1 pt at a time down to the
minimum size; if it still does not fit, a warning is printed: cut the text.
"""

import re
import sys
import uuid
from collections import namedtuple
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import lian_spec as S  # noqa: E402

Box = namedtuple("Box", "x y w h")
TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "lian-deck-template.pptx"
G = S.SLIDE
_SPAN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)")
_SCRIPT = re.compile(r"(\^\{.+?\}|_\{.+?\})")
WARNINGS = []


def _plain_shape(shape):
    """Drop the theme style reference so no shadow or theme colour leaks in."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    return shape


def _rgb(hex_):
    return RGBColor.from_string(hex_)


def runs(text):
    """Split inline markup into (text, bold, italic, baseline) pieces.

    **bold** and *italic* may contain ^{sup} and _{sub}; they do not nest in each other.
    """
    out = []
    for piece in _SPAN.split(text):
        if not piece:
            continue
        bold = piece.startswith("**") and piece.endswith("**") and len(piece) > 4
        italic = not bold and piece.startswith("*") and piece.endswith("*") and len(piece) > 2
        inner = piece[2:-2] if bold else piece[1:-1] if italic else piece
        for sub in _SCRIPT.split(inner):
            if not sub:
                continue
            if sub.startswith("^{"):
                out.append((sub[2:-1], bold, italic, 30000))
            elif sub.startswith("_{"):
                out.append((sub[2:-1], bold, italic, -25000))
            else:
                out.append((sub, bold, italic, 0))
    return out


def _plain(text):
    return "".join(r[0] for r in runs(text))


def _set_font(run, size, bold=False, italic=False, colour=S.BLACK, baseline=0):
    f = run.font
    f.name, f.size, f.bold, f.italic = S.FONT, Pt(size), bold, italic
    f.color.rgb = _rgb(colour)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", S.FONT)
    if baseline:
        rpr.set("baseline", str(baseline))


def _write(p, text, size, bold=False, italic=False, colour=S.BLACK):
    for t, b, i, base in runs(text):
        r = p.add_run()
        r.text = t
        _set_font(r, size, bold or b, italic or i, colour, base)


def _bullet(p, level, numbered, size):
    ppr = p._p.get_or_add_pPr()
    indent = 0.28 if level == 0 else 0.30
    ppr.set("marL", str(int(Inches(indent * (level + 1)))))
    ppr.set("indent", str(int(-Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buSzPct"):
        el = ppr.find(qn(tag))
        if el is not None:
            ppr.remove(el)
    font = etree.SubElement(ppr, qn("a:buFont"))
    font.set("typeface", S.FONT)
    if numbered and level == 0:
        num = etree.SubElement(ppr, qn("a:buAutoNum"))
        num.set("type", "arabicPeriod")
    else:
        ch = etree.SubElement(ppr, qn("a:buChar"))
        ch.set("char", "•" if level == 0 else "–")


def _no_bullet(p):
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", "0")
    ppr.set("indent", "0")
    if ppr.find(qn("a:buNone")) is None:
        etree.SubElement(ppr, qn("a:buNone"))


def _frame(shape, inset=0.0, anchor=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(inset)
    tf.margin_top = tf.margin_bottom = Inches(inset * 0.6)
    return tf


def _lines(text, size, width_in, indent_in=0.0):
    """Rough line count for Times New Roman (average glyph about 0.46 em)."""
    usable = max(width_in - indent_in, 0.5) * 72
    per_line = max(int(usable / (0.46 * size)), 1)
    words, lines, cur = _plain(text).split(), 1, 0
    for w in words:
        need = len(w) + (1 if cur else 0)
        if cur + need > per_line:
            lines, cur = lines + 1, len(w)
        else:
            cur += need
    return lines


def _height(items, size, width_in, bullets, heading=None, head_size=None, spacing=1.15, after_pt=6):
    h = 0.0
    if heading:
        h += _lines(heading, head_size, width_in) * head_size * spacing / 72 + 6 / 72
    for it in items:
        level, text = it if isinstance(it, tuple) else (0, it)
        ind = (0.28 * (level + 1)) if bullets else 0.0
        h += _lines(text, size, width_in, ind) * size * spacing / 72 + after_pt / 72
    return h


class Deck:
    def __init__(self, short_title="", template=TEMPLATE, footer=True):
        self.prs = Presentation(str(template)) if Path(template).exists() else Presentation()
        if not Path(template).exists():
            self.prs.slide_width, self.prs.slide_height = Inches(S.SLIDE_W_IN), Inches(S.SLIDE_H_IN)
        self.short_title = short_title
        self.footer = footer
        self.blank = next((l for l in self.prs.slide_layouts if l.name.lower() == "blank"), self.prs.slide_layouts[-1])
        self.W = self.prs.slide_width / 914400
        self.H = self.prs.slide_height / 914400

    # ------------------------------------------------------------ geometry
    def body(self):
        m = G["margin_x"]
        return Box(m, G["body_y"], self.W - 2 * m, G["body_bottom"] - G["body_y"])

    def columns(self, n=2, box=None, widths=None, gap=None):
        box = box or self.body()
        gap = G["gap"] if gap is None else gap
        widths = widths or [1] * n
        total = box.w - gap * (len(widths) - 1)
        out, x = [], box.x
        for w in widths:
            ww = total * w / sum(widths)
            out.append(Box(x, box.y, ww, box.h))
            x += ww + gap
        return out

    def rows(self, n=2, box=None, heights=None, gap=None):
        box = box or self.body()
        gap = G["gap"] * 0.6 if gap is None else gap
        heights = heights or [1] * n
        total = box.h - gap * (len(heights) - 1)
        out, y = [], box.y
        for h in heights:
            hh = total * h / sum(heights)
            out.append(Box(box.x, y, box.w, hh))
            y += hh + gap
        return out

    def grid(self, r=2, c=2, box=None):
        return [cell for row in self.rows(r, box) for cell in self.columns(c, row)]

    def split(self, box, bottom_h, gap=None):
        """Return (top, bottom): a strip of height bottom_h at the foot of box."""
        gap = G["gap"] * 0.6 if gap is None else gap
        return Box(box.x, box.y, box.w, box.h - bottom_h - gap), Box(box.x, box.y + box.h - bottom_h, box.w, bottom_h)

    # ------------------------------------------------------------ primitives
    def _textbox(self, slide, box, name=None):
        tb = slide.shapes.add_textbox(Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h))
        if name:
            tb.name = name
        return tb

    def _line(self, slide, x, y, w, weight=0.75, colour=S.BLACK):
        ln = _plain_shape(slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y)))
        ln.line.color.rgb = _rgb(colour)
        ln.line.width = Pt(weight)
        return ln

    def _footer(self, slide):
        if not self.footer:
            return
        m, y, h, pt = G["margin_x"], G["footer_y"], G["footer_h"], G["footer_pt"]
        left = self._textbox(slide, Box(m, y, self.W * 0.7, h), "Footer")
        tf = _frame(left, anchor=MSO_ANCHOR.MIDDLE)
        _write(tf.paragraphs[0], self.short_title, pt, colour=S.GREY_TEXT)
        right = self._textbox(slide, Box(self.W - m - 1.0, y, 1.0, h), "Slide number")
        tf = _frame(right, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        fld = etree.SubElement(p._p, qn("a:fld"))
        fld.set("id", "{" + str(uuid.uuid4()).upper() + "}")
        fld.set("type", "slidenum")
        rpr = etree.SubElement(fld, qn("a:rPr"))
        rpr.set("lang", "en-US")
        rpr.set("sz", str(pt * 100))
        fill = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", S.GREY_TEXT)
        for tag in ("a:latin", "a:ea", "a:cs"):
            etree.SubElement(rpr, qn(tag)).set("typeface", S.FONT)
        etree.SubElement(fld, qn("a:t")).text = str(len(self.prs.slides))

    def notes(self, slide, text):
        slide.notes_slide.notes_text_frame.text = text

    # ------------------------------------------------------------ slide types
    def title_slide(self, title, subtitle=(), byline=None, date=None):
        s = self.prs.slides.add_slide(self.blank)
        m = 1.0
        w = self.W - 2 * m
        size = 36 if len(_plain(title)) <= 90 else 32
        th = _lines(title, size, w) * size * 1.15 / 72 + 0.2
        y = max(1.6, 3.2 - th)
        tb = self._textbox(s, Box(m, y, w, th), "Title")
        tf = _frame(tb, anchor=MSO_ANCHOR.BOTTOM)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _write(p, title, size, bold=True)
        y += th + 0.25
        self._line(s, self.W / 2 - 1.5, y, 3.0, weight=1.0)
        y += 0.35
        lines = [(t, 20) for t in subtitle] + ([(byline, 18)] if byline else []) + ([(date, 16)] if date else [])
        if lines:
            sb = self._textbox(s, Box(m, y, w, self.H - y - 0.8), "Subtitle")
            tf = _frame(sb)
            for i, (t, sz) in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                p.space_after = Pt(4 if i < len(subtitle) - 1 else 14)
                _write(p, t, sz)
        return s

    def slide(self, title, key_message=None):
        """A content slide: title, optional one-sentence key message, rule, footer."""
        s = self.prs.slides.add_slide(self.blank)
        m, w = G["margin_x"], self.W - 2 * G["margin_x"]
        size = G["title_pt"] if _lines(title, G["title_pt"], w) == 1 else G["title_pt"] - 4
        tb = self._textbox(s, Box(m, G["title_y"], w, G["title_h"]), "Title")
        tf = _frame(tb, anchor=MSO_ANCHOR.BOTTOM)
        _write(tf.paragraphs[0], title, size, bold=True)
        if key_message:
            kb = self._textbox(s, Box(m, G["key_y"], w, G["key_h"]), "Key message")
            tf = _frame(kb, anchor=MSO_ANCHOR.MIDDLE)
            ksize = G["key_pt"] if _lines(key_message, G["key_pt"], w) == 1 else G["key_pt"] - 2
            if _lines(key_message, ksize, w) > 1:
                WARNINGS.append(f"slide {len(self.prs.slides)}: key message runs past one line; shorten it")
            _write(tf.paragraphs[0], key_message, ksize, italic=True)
        self._line(s, m, G["rule_y"], w)
        self._footer(s)
        return s

    def section_slide(self, title, subtitle=None):
        s = self.prs.slides.add_slide(self.blank)
        m, w = 1.0, self.W - 2.0
        tb = self._textbox(s, Box(m, 2.6, w, 1.0), "Title")
        tf = _frame(tb, anchor=MSO_ANCHOR.BOTTOM)
        _write(tf.paragraphs[0], title, 34, bold=True)
        self._line(s, m, 3.75, 3.0, weight=1.0)
        if subtitle:
            sb = self._textbox(s, Box(m, 3.95, w, 0.8), "Subtitle")
            _write(_frame(sb).paragraphs[0], subtitle, 20, italic=True)
        self._footer(s)
        return s

    def takeaway(self, headline, items=(), kicker="Takeaway"):
        s = self.prs.slides.add_slide(self.blank)
        m, w = 1.0, self.W - 2.0
        kb = self._textbox(s, Box(m, 1.0, w, 0.5), "Kicker")
        _write(_frame(kb).paragraphs[0], kicker.upper(), 16, bold=True)
        hsize = 32
        hh = _lines(headline, hsize, w) * hsize * 1.15 / 72 + 0.15
        hb = self._textbox(s, Box(m, 1.55, w, hh), "Title")
        _write(_frame(hb).paragraphs[0], headline, hsize, bold=True)
        y = 1.55 + hh + 0.2
        self._line(s, m, y, w)
        if items:
            self.text(s, Box(m, y + 0.3, w, G["body_bottom"] - y - 0.3), items, size=22)
        self._footer(s)
        return s

    def closing(self, title="Thank you", subtitle="Questions and discussion"):
        s = self.prs.slides.add_slide(self.blank)
        tb = self._textbox(s, Box(1.0, 2.6, self.W - 2, 1.0), "Title")
        tf = _frame(tb, anchor=MSO_ANCHOR.BOTTOM)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        _write(tf.paragraphs[0], title, 40, bold=True)
        self._line(s, self.W / 2 - 1.5, 3.8, 3.0, weight=1.0)
        if subtitle:
            sb = self._textbox(s, Box(1.0, 4.0, self.W - 2, 0.7), "Subtitle")
            tf = _frame(sb)
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            _write(tf.paragraphs[0], subtitle, 22, italic=True)
        return s

    # ------------------------------------------------------------ body content
    def text(self, slide, box, items, heading=None, size=None, bullets=True, numbered=False,
             inset=0.0, name=None):
        """Bulleted (or plain, or numbered) text with an optional bold heading."""
        items = list(items)
        start = size or G["body_pt"]
        head_size = G["panel_head_pt"]
        size = start
        while size > G["min_pt"] and _height(items, size, box.w - 2 * inset, bullets, heading, head_size) > box.h - 1.2 * inset:
            size -= 1
        if _height(items, size, box.w - 2 * inset, bullets, heading, head_size) > box.h - 1.2 * inset:
            WARNINGS.append(f"slide {len(self.prs.slides)}: text in '{name or heading or 'text'}' may overflow at {size} pt; cut it")
        tb = self._textbox(slide, box, name or heading or "Text")
        tf = _frame(tb, inset)
        first = True
        if heading:
            p = tf.paragraphs[0]
            _no_bullet(p)
            p.space_after = Pt(6)
            _write(p, heading, head_size, bold=True)
            first = False
        for it in items:
            level, t = it if isinstance(it, tuple) else (0, it)
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(6)
            p.line_spacing = 1.0
            if bullets:
                _bullet(p, level, numbered, size)
            else:
                _no_bullet(p)
            _write(p, t, size if level == 0 else max(size - 2, G["min_pt"]))
        return tb

    def panel(self, slide, box, items, heading=None, size=None, bullets=True, numbered=False):
        """Light-grey panel (the 'card' on the reference deck), no outline."""
        r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h))
        _plain_shape(r)
        r.name = (heading or "Panel") + " box"
        r.fill.solid()
        r.fill.fore_color.rgb = _rgb(S.GREY_FILL)
        r.line.fill.background()
        return self.text(slide, box, items, heading=heading, size=size or G["body_pt"], bullets=bullets,
                         numbered=numbered, inset=0.25, name=heading)

    def callout(self, slide, box, label, text, size=None):
        """Boxed statement: thin black outline, bold label. One per slide at most."""
        r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h))
        _plain_shape(r)
        r.name = "Callout"
        r.fill.background()
        r.line.color.rgb = _rgb(S.BLACK)
        r.line.width = Pt(1.0)
        tf = _frame(r, 0.25, MSO_ANCHOR.MIDDLE)
        size = size or G["body_pt"]
        full = f"**{label}** {text}" if label else text
        while size > G["min_pt"] and _height([full], size, box.w - 0.5, False) > box.h - 0.1:
            size -= 1
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        _no_bullet(p)
        _write(p, full, size)
        return r

    def table(self, slide, box, rows, col_widths=None, size=None, header=True, first_col_bold=False,
              align=None):
        """Three-line table: rule above, rule under the header, rule below. No fills, no grid.

        rows: list of lists of strings (markup allowed); first row is the header.
        align: optional list like ["l", "r", "r"] per column.
        """
        n, k = len(rows), len(rows[0])
        size = size or G["table_pt"]
        col_widths = col_widths or [1] * k
        widths = [box.w * c / sum(col_widths) for c in col_widths]
        row_h = min(box.h / n, (size * 1.25 + 8) / 72 * 1.0 + 0.06)
        gs = slide.shapes.add_table(n, k, Inches(box.x), Inches(box.y), Inches(box.w), Inches(row_h * n))
        gs.name = "Table"
        tbl = gs.table
        tblpr = tbl._tbl.tblPr
        for a in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
            tblpr.set(a, "0")
        sid = tblpr.find(qn("a:tableStyleId"))
        if sid is None:
            sid = etree.SubElement(tblpr, qn("a:tableStyleId"))
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
        for j, w in enumerate(widths):
            tbl.columns[j].width = Inches(w)
        for i in range(n):
            tbl.rows[i].height = Inches(row_h)
            for j in range(k):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                a = (align[j] if align else ("l" if j == 0 else "l")).lower()
                p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[a]
                bold = (header and i == 0) or (first_col_bold and j == 0 and i > 0)
                _write(p, str(rows[i][j]), size, bold=bold)
                top = 1.5 if i == 0 else (0.75 if header and i == 1 else 0)
                bottom = 1.5 if i == n - 1 else (0.75 if header and i == 0 else 0)
                _cell_borders(cell, top, bottom)
        est = sum(max(_lines(str(rows[i][j]), size, widths[j] - 0.16) for j in range(k)) * size * 1.25 / 72 + 0.1
                  for i in range(n))
        if est > box.h + 0.05:
            WARNINGS.append(f"slide {len(self.prs.slides)}: table is about {est:.1f} in tall for a {box.h:.1f} in box; "
                            "cut rows or text, or lower size")
        return gs

    def image(self, slide, box, path, caption=None, caption_pt=12):
        from PIL import Image
        cap_h = 0.4 if caption else 0.0
        with Image.open(path) as im:
            iw, ih = im.size
        avail_w, avail_h = box.w, box.h - cap_h
        scale = min(avail_w / iw, avail_h / ih)
        w, h = iw * scale, ih * scale
        x, y = box.x + (box.w - w) / 2, box.y + (avail_h - h) / 2
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        if caption:
            cb = self._textbox(slide, Box(box.x, y + h + 0.05, box.w, cap_h), "Caption")
            p = _frame(cb).paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _write(p, caption, caption_pt)
        return pic

    def source(self, slide, text):
        """Small source line just above the footer."""
        m = G["margin_x"]
        sb = self._textbox(slide, Box(m, G["body_bottom"] + 0.02, self.W - 2 * m, 0.22), "Source")
        _write(_frame(sb).paragraphs[0], text, G["min_pt"], colour=S.GREY_TEXT)

    def save(self, path):
        self.prs.save(str(path))
        for w in WARNINGS:
            print("WARNING:", w)
        WARNINGS.clear()
        print("wrote", path, f"({len(self.prs.slides)} slides)")


def _cell_borders(cell, top_pt, bottom_pt):
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:noFill", "a:solidFill"):
        for el in tcpr.findall(qn(tag)):
            tcpr.remove(el)
    for tag, pt in (("a:lnL", 0), ("a:lnR", 0), ("a:lnT", top_pt), ("a:lnB", bottom_pt)):
        ln = etree.SubElement(tcpr, qn(tag))
        if pt:
            ln.set("w", str(int(pt * 12700)))
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", S.BLACK)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
    etree.SubElement(tcpr, qn("a:noFill"))


# ---------------------------------------------------------------- template
def build_template(out=TEMPLATE):
    """Write assets/lian-deck-template.pptx: 16:9, Times New Roman theme fonts,
    black-and-white colour scheme, layouts rescaled to the wide page."""
    prs = Presentation()
    old_w = prs.slide_width
    prs.slide_width, prs.slide_height = Inches(S.SLIDE_W_IN), Inches(S.SLIDE_H_IN)
    fx = prs.slide_width / old_w
    for shapes in [prs.slide_master.shapes] + [l.shapes for l in prs.slide_layouts]:
        for sh in shapes:
            if sh.left is not None:
                sh.left, sh.width = int(sh.left * fx), int(sh.width * fx)
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ns = {"a": a}
    theme = next(p for p in prs.slide_master.part.package.iter_parts() if str(p.partname).startswith("/ppt/theme/"))
    root = etree.fromstring(theme.blob)
    for tag in ("majorFont", "minorFont"):
        for latin in root.iterfind(f".//a:{tag}/a:latin", ns):
            latin.set("typeface", S.FONT)
    scheme = {"dk1": S.BLACK, "lt1": S.WHITE, "dk2": "404040", "lt2": S.GREY_FILL, "accent1": S.BLACK,
              "accent2": "595959", "accent3": "7F7F7F", "accent4": "A6A6A6", "accent5": "D9D9D9",
              "accent6": S.RED, "hlink": S.BLACK, "folHlink": "404040"}
    cs = root.find(".//a:clrScheme", ns)
    cs.set("name", "Lian black and white")
    for key, val in scheme.items():
        el = cs.find(f"a:{key}", ns)
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{a}}}srgbClr").set("val", val)
    root.find(".//a:fontScheme", ns).set("name", "Lian Times New Roman")
    theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    prs.core_properties.title = "Lian signature style: deck template"
    prs.core_properties.author = ""
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print("wrote", out)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "template":
        build_template(Path(sys.argv[2]) if len(sys.argv) > 2 else TEMPLATE)
    else:
        print(__doc__)
