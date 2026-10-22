"""Word side of the Lian signature style: build the template, finish any .docx.

  python docx_style.py template [--paper letter|a4] [--out-dir ../assets]
      Writes lian-article-reference.docx (pandoc --reference-doc, python-docx
      base) and lian-article-template.dotx (double-click in Word for a new
      document). Needs pandoc for the full set of pandoc style names; without
      pandoc it starts from python-docx's default document.

  python docx_style.py finish in.docx [out.docx] [--paper letter|a4]
      [--spacing double|onehalf|single] [--allow-red] [--keep-colour]
      Enforces the style on a .docx built by any tool (pandoc, docx-js,
      python-docx, Word): Times New Roman everywhere, black text, 1-inch
      margins, page numbers, three-line tables with 10 pt single-spaced text.
      Without out.docx the file is changed in place.
"""

import argparse
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent))
import lian_spec as S  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
LINE = {"double": 2.0, "onehalf": 1.5, "single": 1.0}
BODY_STYLES = ("Normal", "Body Text", "First Paragraph")
THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


# ------------------------------------------------------------------ helpers
def _set_rfonts(rpr, font=S.FONT):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in THEME_ATTRS:
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), font)


def _font(style, size=S.BODY_PT, bold=False, italic=False, font=S.FONT):
    f = style.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = RGBColor.from_string(S.BLACK)
    _set_rfonts(style.element.get_or_add_rPr(), font)
    col = style.element.rPr.find(qn("w:color"))
    if col is not None:
        for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
            col.attrib.pop(qn(a), None)


def _para(style, line=1.0, before=0, after=0, first=0.0, left=0.0, right=0.0,
          align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=False):
    p = style.paragraph_format
    p.line_spacing, p.space_before, p.space_after = line, Pt(before), Pt(after)
    p.first_line_indent, p.left_indent, p.right_indent = Inches(first), Inches(left), Inches(right)
    p.alignment, p.keep_with_next = align, keep_next
    p.widow_control = True


def _find_style(doc, name):
    """Look a style up by its display name (pandoc capitalises built-in names,
    which python-docx's own lookup does not expect)."""
    for st in doc.styles:
        if st.name == name or st.style_id == name.replace(" ", ""):
            return st
    return None


def _get_style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH, base="Normal"):
    st = _find_style(doc, name)
    if st is None:
        st = doc.styles.add_style(name, kind)
        if base and kind == WD_STYLE_TYPE.PARAGRAPH:
            st.base_style = _find_style(doc, base)
    return st


def _page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        if "PAGE" in p._p.xml:
            return
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent, pf.line_spacing, pf.space_after = Inches(0), 1.0, Pt(0)
    run = p.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text.strip() == "PAGE":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)
    run.font.name, run.font.size = S.FONT, Pt(S.BODY_PT)
    _set_rfonts(run._r.get_or_add_rPr())


def _page_setup(doc, paper):
    w, h = S.PAPER[paper]
    for s in doc.sections:
        s.page_width, s.page_height = Inches(w), Inches(h)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(S.MARGIN_IN)
        s.header_distance = s.footer_distance = Inches(0.5)
        s.gutter = Inches(0)
        _page_number_footer(s)


# ------------------------------------------------------------------ template
def build_styles(doc, spacing="double"):
    body = LINE[spacing]
    # document defaults: Times New Roman, 12 pt
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            _set_rfonts(rpr)
            for tag in ("w:sz", "w:szCs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = OxmlElement(tag)
                    _insert_rpr(rpr, el)
                el.set(qn("w:val"), str(S.BODY_PT * 2))

    normal = _find_style(doc, "Normal")
    _font(normal)
    _para(normal, line=body, first=S.FIRST_LINE_INDENT_IN)

    for name, first in (("Body Text", S.FIRST_LINE_INDENT_IN), ("First Paragraph", 0.0)):
        st = _get_style(doc, name)
        _font(st)
        _para(st, line=body, first=first)

    compact = _get_style(doc, "Compact")
    _font(compact)
    _para(compact, line=1.0)

    title = _get_style(doc, "Title")
    _font(title, S.TITLE_PT, bold=True)
    _para(title, align=WD_ALIGN_PARAGRAPH.CENTER, after=12, keep_next=True)
    for name, after, italic in (("Subtitle", 12, False), ("Author", 0, False), ("Date", 24, False)):
        st = _get_style(doc, name)
        _font(st, italic=italic)
        _para(st, align=WD_ALIGN_PARAGRAPH.CENTER, after=after, keep_next=True)

    at = _get_style(doc, "Abstract Title")
    _font(at, bold=True)
    _para(at, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6, keep_next=True)
    ab = _get_style(doc, "Abstract")
    _font(ab)
    _para(ab, left=0.5, right=0.5, after=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # headings: 1 bold, 2 bold italic, 3+ italic; all 12 pt black, flush left
    for i in range(1, 10):
        st = _get_style(doc, f"Heading {i}")
        _font(st, bold=i <= 2, italic=i >= 2)
        _para(st, line=body, before=12, after=0, keep_next=True)

    bt = _get_style(doc, "Block Text")
    _font(bt)
    _para(bt, left=0.5, right=0.5, before=6, after=12)

    fn = _get_style(doc, "Footnote Text")
    _font(fn, S.SMALL_PT)
    _para(fn, after=2)

    tc = _get_style(doc, "Table Caption")
    _font(tc)
    _para(tc, before=12, after=6, keep_next=True)
    for name in ("Image Caption", "Caption"):
        st = _get_style(doc, name)
        _font(st)
        _para(st, before=6, after=12)
    for name in ("Figure", "Captioned Figure"):
        st = _get_style(doc, name)
        _font(st)
        _para(st, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, keep_next=True)

    # notes and sources under tables and figures: ::: {custom-style="Table Note"}
    for name in ("Table Note", "Figure Note"):
        st = _get_style(doc, name)
        _font(st, S.SMALL_PT)
        _para(st, after=4)

    bib = _get_style(doc, "Bibliography")
    _font(bib)
    _para(bib, first=-S.HANGING_INDENT_IN, left=S.HANGING_INDENT_IN, after=6)

    for name in ("Definition Term",):
        st = _get_style(doc, name)
        _font(st, bold=True)
        _para(st, keep_next=True)
    d = _get_style(doc, "Definition")
    _font(d)
    _para(d, left=0.5, after=6)
    toc = _get_style(doc, "TOC Heading")
    _font(toc, bold=True)
    _para(toc, before=12, after=6)

    # character styles
    for name, kw in (("Hyperlink", {}), ("Footnote Reference", {}), ("Verbatim Char", {"font": S.FONT_CODE, "size": S.SMALL_PT})):
        st = _find_style(doc, name) or doc.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
        _font(st, size=kw.get("size", S.BODY_PT), font=kw.get("font", S.FONT))
        st.font.underline = False
        if name == "Footnote Reference":
            st.font.superscript = True
    _three_line_table_style(doc)


def _border(tag, sz):
    el = OxmlElement(tag)
    if sz:
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), S.BLACK)
    else:
        el.set(qn("w:val"), "nil")
    return el


def _three_line_table_style(doc):
    """Pandoc applies the table style named 'Table'. Give it booktabs rules."""
    st = _find_style(doc, "Table") or doc.styles.add_style("Table", WD_STYLE_TYPE.TABLE)
    el = st.element
    tblpr = el.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        el.append(tblpr)
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    b = OxmlElement("w:tblBorders")
    for tag, sz in (("w:top", 12), ("w:left", 0), ("w:bottom", 12), ("w:right", 0), ("w:insideH", 0), ("w:insideV", 0)):
        b.append(_border(tag, sz))
    tblpr.append(b)
    for old in el.findall(qn("w:tblStylePr")):
        el.remove(old)
    fr = OxmlElement("w:tblStylePr")
    fr.set(qn("w:type"), "firstRow")
    tcpr = OxmlElement("w:tcPr")
    tcb = OxmlElement("w:tcBorders")
    tcb.append(_border("w:bottom", 6))
    tcpr.append(tcb)
    fr.append(tcpr)
    el.append(fr)


def _clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _theme_fonts(path):
    """Point the theme's major and minor fonts at Times New Roman."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/theme/") and item.filename.endswith(".xml"):
                root = etree.fromstring(data)
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for tag in ("majorFont", "minorFont"):
                    for latin in root.iterfind(f".//a:{tag}/a:latin", ns):
                        latin.set("typeface", S.FONT)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    tmp.replace(path)


def _as_dotx(src, dst):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
            zout.writestr(item, data)


def build_template(out_dir, paper=S.DEFAULT_PAPER):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    ref = out_dir / "lian-article-reference.docx"
    base = None
    if shutil.which("pandoc"):
        base = Path(tempfile.mkdtemp()) / "pandoc-ref.docx"
        subprocess.run(["pandoc", "-o", str(base), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(str(base)) if base else Document()
    _clear_body(doc)
    build_styles(doc)
    _page_setup(doc, paper)
    doc.core_properties.title = "Lian signature style: article template"
    doc.core_properties.author = ""
    doc.save(str(ref))
    _theme_fonts(ref)
    _as_dotx(ref, out_dir / "lian-article-template.dotx")
    print("wrote", ref, "and", out_dir / "lian-article-template.dotx")


# ------------------------------------------------------------------ finish
def _xml_parts(doc):
    """Yield (root, commit) for every XML part that can carry text formatting."""
    for part in doc.part.package.iter_parts():
        name = str(part.partname)
        if not (name.startswith("/word/") and name.endswith(".xml")):
            continue
        if any(k in name for k in ("document", "styles", "footnotes", "endnotes", "header", "footer", "numbering", "comments")):
            if hasattr(part, "_element"):
                yield part._element, (lambda: None)
            else:
                root = etree.fromstring(part.blob)

                def commit(part=part, root=root):
                    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

                yield root, commit


def _fix_fonts_and_colours(root, allow_red, keep_colour):
    for rf in root.iter(qn("w:rFonts")):
        fonts = {rf.get(qn(a)) for a in ("w:ascii", "w:hAnsi")}
        is_code = bool(fonts & {"Courier New", "Consolas", "Menlo", "Courier"})
        for a in THEME_ATTRS:
            rf.attrib.pop(qn(a), None)
        target = S.FONT_CODE if is_code else S.FONT
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), target)
    if keep_colour:
        return
    reds = {"C00000", "FF0000"}
    for col in list(root.iter(qn("w:color"))):
        val = (col.get(qn("w:val")) or "").upper()
        if allow_red and val in reds:
            continue
        col.set(qn("w:val"), "auto" if val in ("AUTO", "") else S.BLACK)
        for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
            col.attrib.pop(qn(a), None)
    for hl in list(root.iter(qn("w:highlight"))):
        if not (allow_red and hl.get(qn("w:val")) == "red"):
            hl.getparent().remove(hl)


def _three_line(tbl, table_pt):
    tblpr = tbl.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        tbl.insert(0, tblpr)
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    b = OxmlElement("w:tblBorders")
    for tag, sz in (("w:top", 12), ("w:left", 0), ("w:bottom", 12), ("w:right", 0), ("w:insideH", 0), ("w:insideV", 0)):
        b.append(_border(tag, sz))
    # schema order: tblBorders comes after tblW/jc/tblCellSpacing/tblInd, before shd/tblLayout/tblCellMar/tblLook
    anchor = None
    for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"):
        anchor = tblpr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(b)
    else:
        tblpr.append(b)
    shd = tblpr.find(qn("w:shd"))
    if shd is not None:
        tblpr.remove(shd)
    rows = tbl.findall(qn("w:tr"))
    for i, tr in enumerate(rows):
        for tc in tr.findall(qn("w:tc")):
            tcpr = tc.find(qn("w:tcPr"))
            if tcpr is None:
                tcpr = OxmlElement("w:tcPr")
                tc.insert(0, tcpr)
            for tag in ("w:tcBorders", "w:shd"):
                el = tcpr.find(qn(tag))
                if el is not None:
                    tcpr.remove(el)
            if i == 0 and len(rows) > 1:
                tcb = OxmlElement("w:tcBorders")
                tcb.append(_border("w:bottom", 6))
                # tcBorders sits after tcW/gridSpan/vMerge
                after = [tcpr.find(qn(t)) for t in ("w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge")]
                after = [a for a in after if a is not None]
                if after:
                    after[-1].addnext(tcb)
                else:
                    tcpr.insert(0, tcb)
            for p in tc.iter(qn("w:p")):
                ppr = p.find(qn("w:pPr"))
                if ppr is None:
                    ppr = OxmlElement("w:pPr")
                    p.insert(0, ppr)
                for tag in ("w:spacing", "w:ind"):
                    el = ppr.find(qn(tag))
                    if el is not None:
                        ppr.remove(el)
                sp = OxmlElement("w:spacing")
                sp.set(qn("w:before"), "0")
                sp.set(qn("w:after"), "0")
                sp.set(qn("w:line"), "240")
                sp.set(qn("w:lineRule"), "auto")
                ind = OxmlElement("w:ind")
                ind.set(qn("w:firstLine"), "0")
                ind.set(qn("w:left"), "0")
                # keep pPr child order valid: pStyle, keepNext.., spacing, ind, ..., jc
                ref = None
                for tag in ("w:contextualSpacing", "w:jc", "w:textDirection", "w:textAlignment", "w:outlineLvl", "w:rPr", "w:sectPr", "w:pPrChange"):
                    ref = ppr.find(qn(tag))
                    if ref is not None:
                        break
                if ref is not None:
                    ref.addprevious(sp)
                    ref.addprevious(ind)
                else:
                    ppr.append(sp)
                    ppr.append(ind)
                for r in p.iter(qn("w:r")):
                    rpr = r.find(qn("w:rPr"))
                    if rpr is None:
                        rpr = OxmlElement("w:rPr")
                        r.insert(0, rpr)
                    for tag in ("w:sz", "w:szCs"):
                        el = rpr.find(qn(tag))
                        if el is not None:
                            rpr.remove(el)
                    for tag in ("w:sz", "w:szCs"):
                        el = OxmlElement(tag)
                        el.set(qn("w:val"), str(int(table_pt * 2)))
                        _insert_rpr(rpr, el)


_RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
              "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
              "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def _insert_rpr(rpr, el):
    name = etree.QName(el).localname
    idx = _RPR_ORDER.index(name)
    for child in rpr:
        cname = etree.QName(child).localname
        if cname in _RPR_ORDER and _RPR_ORDER.index(cname) > idx:
            child.addprevious(el)
            return
    rpr.append(el)


def finish(src, dst=None, paper=S.DEFAULT_PAPER, spacing=None, allow_red=False, keep_colour=False,
           table_pt=S.SMALL_PT):
    dst = dst or src
    doc = Document(str(src))
    _page_setup(doc, paper)
    build_styles(doc, spacing or "double")
    for root, commit in _xml_parts(doc):
        _fix_fonts_and_colours(root, allow_red, keep_colour)
        commit()
    for tbl in doc.element.body.iter(qn("w:tbl")):
        _three_line(tbl, table_pt)
    doc.save(str(dst))
    print("finished", dst)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("template")
    t.add_argument("--paper", default=S.DEFAULT_PAPER, choices=S.PAPER)
    t.add_argument("--out-dir", default=str(Path(__file__).resolve().parent.parent / "assets"))
    f = sub.add_parser("finish")
    f.add_argument("src")
    f.add_argument("dst", nargs="?")
    f.add_argument("--paper", default=S.DEFAULT_PAPER, choices=S.PAPER)
    f.add_argument("--spacing", choices=LINE)
    f.add_argument("--allow-red", action="store_true")
    f.add_argument("--keep-colour", action="store_true")
    f.add_argument("--table-pt", type=float, default=S.SMALL_PT)
    a = ap.parse_args()
    if a.cmd == "template":
        build_template(a.out_dir, a.paper)
    else:
        finish(a.src, a.dst, a.paper, a.spacing, a.allow_red, a.keep_colour, a.table_pt)


if __name__ == "__main__":
    main()
