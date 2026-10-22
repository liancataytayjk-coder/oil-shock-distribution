"""Write the codebook spreadsheet: codebook/codebook.xlsx.

Sheets: README, Datasets, Raw_variables, Clean_variables, Sources, Regions, Results_map.
Every clean variable gets a name, label, units, source, transformation, coverage and the
research question that uses it. Coverage and counts are computed from the data, not typed.
"""

import re

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from catalog import RESULTS, SOURCES
from config import CLEAN, RAW, REGION_NAMES, ROOT

DIV = {"01": "Food and non-alcoholic beverages", "02": "Alcoholic beverages and tobacco", "03": "Clothing and footwear",
       "04": "Housing, water, electricity, gas and other fuels", "05": "Furnishings and household equipment",
       "06": "Health", "07": "Transport", "08": "Information and communication", "09": "Recreation, sport and culture",
       "10": "Education services", "11": "Restaurants and accommodation services",
       "12": "Insurance and financial services", "13": "Personal care and miscellaneous"}
SUB = {"e045": ("04.5", "Electricity, gas and other fuels"), "e0452": ("04.5.2", "Gas (LPG)"),
       "e0453": ("04.5.3", "Liquid fuels (kerosene)"), "e0454": ("04.5.4", "Solid fuels (charcoal, wood)"),
       "t073": ("07.3", "Passenger transport services (fares)"), "f0111": ("01.1.1", "Cereals and cereal products")}
BASKET = {"all": "all income households", "b30": "bottom 30% income households"}
SPLICE = ("old-code PSA back series and current series joined on 2018 = 100; Jan-Aug 2026 chained from "
          "new-code monthly changes")


def describe(name):
    """(label, units, source, transformation, research question) for a clean variable."""
    idx = ("index, 2018 = 100", )
    if name in ("date", "date_str"):
        return ("Month (first day)" if name == "date" else "Month, YYYY-MM text", "date", "-", "-", "all")
    if name in ("year", "month"):
        return (name.capitalize(), "number", "-", "-", "all")
    if name == "t":
        return ("Month counter (Jan 1994 = 0), used for the linear trend", "months", "-", "-", "RQ1-RQ3")
    if name == "region":
        return ("Region code (old geographic code)", "code", "S1-S4", "PSA labels mapped to codes; see Regions sheet", "RQ2")
    if name == "region_name":
        return ("Region name", "text", "S1-S4", "-", "RQ2")
    if name == "new_geo_ext":
        return ("1 if the month is Jan-Aug 2026 and the region's boundary changed in 2024 (VI, VII, XII, BARMM)",
                "0/1", "S1-S4", "Decision D6", "Robustness (appendix)")
    if name.startswith("d_") and name != "d_fuel_nat":
        base = describe(name[2:])
        unit = "percentage points" if name == "d_policy_rate" else "100 x change in ln (approx. % change)"
        tr = "first difference" if name == "d_policy_rate" else "100 x [ln(x_t) - ln(x_t-1)]"
        lab = "Monthly change in the gap ln(CPI all / CPI b30) = d_cpi_all - d_cpi_b30" if name == "d_ratio" \
            else f"Monthly change: {base[0]}"
        rq = base[4]
        if name == "d_cpi_all":
            rq = "RQ1 outcome (pi^A); RQ3"
        if name == "d_cpi_b30":
            rq = "RQ1 outcome (pi^B); RQ3"
        if name == "d_ratio":
            rq = "RQ2 outcome (Delta D)"
        if name == "d_fuel":
            lab, rq = "Fuel price shock f_t: monthly change in the fuel index 07.2.2", "Shock in RQ1-RQ3"
        return (lab, unit, base[2], tr, rq)
    if name == "d_fuel_nat":
        return ("National fuel price shock, merged into the regional panel", "100 x change in ln", "S1-S2",
                "national d_fuel", "Robustness (appendix)")
    if name == "cpi_all":
        return ("CPI, all income households, all items", *idx, "S1-S2", SPLICE, "RQ1-RQ3")
    if name == "cpi_b30":
        return ("CPI, bottom 30% income households, all items", *idx, "S3-S4",
                "PSA backcasts 2000-2011 and 2012-2017 joined to the 2018 series (D7); " + SPLICE.split(";")[1].strip(),
                "RQ1-RQ3")
    if name == "fuel":
        return ("Retail fuel index: CPI 07.2.2 fuels and lubricants for personal transport (all income)", *idx,
                "S1-S2", SPLICE, "Shock")
    if name == "gasoline":
        return ("CPI 07.2.2.2 gasoline (all income), from 2018", *idx, "S2", "-", "Robustness (appendix)")
    if name == "ratio":
        return ("Inflation-gap level D = 100 x ln(CPI all / CPI b30)", "100 x ln ratio", "S1-S4",
                "computed from cpi_all and cpi_b30", "Objective 1; RQ2")
    m = re.match(r"^(b30_)?c(\d{2})$", name)
    if m:
        b = "b30" if m.group(1) else "all"
        return (f"CPI division {m.group(2)} {DIV[m.group(2)]}, {BASKET[b]}", *idx,
                "S3-S4" if b == "b30" else "S1-S2",
                SPLICE if b == "all" else "from Jan 2012 (back series); " + SPLICE.split(";")[1].strip(), "RQ2 why")
    m = re.match(r"^(all|b30)_(\w+)$", name)
    if m and m.group(2) in SUB:
        code, lab = SUB[m.group(2)]
        return (f"CPI {code} {lab}, {BASKET[m.group(1)]}", *idx, "S1-S2" if m.group(1) == "all" else "S3-S4",
                SPLICE, "RQ3")
    macro = {
        "brent_usd": ("Brent crude oil price", "USD per barrel", "S6", "extract of the Pink Sheet", "-"),
        "dubai_usd": ("Dubai crude oil price", "USD per barrel", "S6", "extract of the Pink Sheet", "Robustness"),
        "php_usd": ("Exchange rate, pesos per US dollar, monthly average", "PHP per USD", "S8", "extract", "Robustness"),
        "dubai_php": ("Dubai crude oil price in pesos", "PHP per barrel", "S6, S8", "dubai_usd x php_usd", "Robustness"),
        "wb_gasoline_php": ("Pump price, Manila regular gasoline (RON 91)", "PHP per litre", "S7", "extract",
                            "Fuel-index check; robustness (appendix)"),
        "policy_rate": ("BSP policy rate, end of month", "percent", "S9", "-", "Robustness"),
        "ffpi": ("FAO Food Price Index, nominal", "index, 2014-2016 = 100", "S10", "-", "Robustness"),
        "pos": ("1 if the fuel index rose this month", "0/1", "S1-S2", "d_fuel > 0", "Robustness (appendix)"),
        "jan2012": ("1 in January 2012 (base-change join month)", "0/1", "-", "Decision D7", "Robustness (appendix)"),
        "jan2018": ("1 in January 2018 (base-change join month)", "0/1", "-", "Decision D7", "Robustness (appendix)"),
        "poverty_2018": ("Poverty incidence among population, 2018", "percent", "S5", "-", "Region groups"),
        "poverty_2023": ("Poverty incidence among population, 2023", "percent", "S5", "-", "Descriptive"),
        "threshold_2023": ("Annual per-capita poverty threshold, 2023", "PHP per person per year", "S5", "-",
                           "RQ3 peso cost"),
        "food_share_2018": ("Food share of the region's own all-income basket, 2018", "percent of basket", "S1",
                            "100 x weight(01) / weight(all items)", "Region groups"),
        "high_poverty": ("1 if 2018 poverty incidence is above the median of the 17 regions", "0/1", "S5",
                         "defined before estimation (D8)", "Appendix (region groups)"),
        "high_food": ("1 if the 2018 food share is above the median of the 17 regions", "0/1", "S1",
                      "defined before estimation (D8)", "Appendix (region groups)"),
        "mindanao": ("1 for Regions IX, X, XI, XII, XIII and BARMM", "0/1", "-", "defined before estimation (D8)",
                     "Appendix (region groups)"),
        "island": ("1 for MIMAROPA, VI, VII and VIII", "0/1", "-", "defined before estimation (D8)",
                   "Appendix (region groups)"),
        "basket": ("CPI basket: all = all income households, b30 = bottom 30%", "code", "S1-S4", "-", "RQ2-RQ3"),
        "geo_code": ("PSA geographic code of the source table: old (17 regions) or new (with NIR)", "code",
                     "S1-S4", "-", "RQ2-RQ3"),
        "coicop": ("COICOP 2018 item code (0 = all items)", "code", "S1-S4", "split from the PSA label", "RQ2-RQ3"),
        "item": ("Item description, as published", "text", "S1-S4", "split from the PSA label", "RQ2-RQ3"),
        "weight": ("Basket weight, 2018", "percent of the national (or regional) basket", "S1-S4", "-", "RQ2-RQ3"),
        "value": ("Index value", "index, 2018 = 100", "S2, S4", "-", "RQ3 decomposition"),
    }
    return macro.get(name, ("(see source)", "-", "-", "-", "-"))


def coverage(df, col):
    if "date" not in df or col in ("date", "date_str"):
        return "", ""
    d = df.loc[df[col].notna(), "date"]
    return (str(d.min())[:7], str(d.max())[:7]) if len(d) else ("", "")


def autosize(ws, widths=None):
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), (widths or {}).get(i, 60))
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def main():
    files = {
        "national_monthly.csv": "One row per month, Jan 1994-Aug 2026: national price indices, macro series and "
                                "their monthly changes (national models, RQ1-RQ3)",
        "regional_panel.csv": "One row per region and month: 17 regions x 392 months, Jan 1994-Aug 2026 "
                              "(balanced; regional panel, RQ2-RQ3)",
        "weights_2018.csv": "2018 CPI basket weights, both baskets, both geographic codes, national and regional",
        "region_info.csv": "One row per region (and the Philippines): poverty statistics, food share, analysis groups",
        "items_2026.csv": "National indices, new geographic code, Dec 2025-Aug 2026, for the 2026 decomposition",
    }
    data = {f: pd.read_csv(CLEAN / f, parse_dates=["date"] if f not in ("weights_2018.csv", "region_info.csv") else None,
                           dtype={"coicop": str}) for f in files}

    readme = pd.DataFrame({"Item": [
        "Study", "Author", "Package", "How to run (Python)", "How to run (Stata 17)", "Raw data", "Clean data",
        "Missing values", "Dates", "Units", "Region codes", "Contact"], "Description": [
        "Who Pays for Oil Shocks? Fuel Price Pass-Through to the Inflation of Poor Households across Philippine Regions",
        "Lian Cataytay, Graduate School of Applied Economics, University of Southeastern Philippines",
        "data/raw (as downloaded, with data/raw/extracts for Excel sheets), data/clean (analysis-ready CSV), "
        "code/python and code/stata (full pipeline), output/ (tables and figures), paper/ (both papers)",
        "From the repository root: ./run_all.sh  (Python 3.11; packages in requirements.txt)",
        "Open code/stata/00_master.do, set the global ROOT to the repository folder, run it. It installs xtscc "
        "and estout from SSC on first run, rebuilds the clean data from data/raw, runs every model and compares "
        "the key results with code/stata/expected_results.csv",
        "Untouched downloads (PSA tables as CSV; World Bank and BSP Excel files with CSV extracts of the sheets used). "
        "See the Sources and Raw_variables sheets.",
        "Built from the raw files by code/python/02_build_clean.py (and code/stata/02_build_clean.do). "
        "See the Datasets and Clean_variables sheets.",
        "Empty cells. Series start at different months (see first/last month in Clean_variables); bottom-30% "
        "divisions start Jan 2012; gasoline Jan 2018.",
        "date is the first day of the month (YYYY-MM-DD); date_str is YYYY-MM. In Stata: gen m = monthly(date_str, \"YM\")",
        "Indices are 2018 = 100. Changes (d_ prefix) are 100 x the change in the natural log: approximately the "
        "percentage change, so coefficients read as percentage points per 1% fuel price change.",
        "PSA old geographic code (17 regions; Negros in VI and VII). See the Regions sheet.",
        "liancataytayjk@gmail.com"]})

    ds = []
    for f, desc in files.items():
        df = data[f]
        ds.append({"File": f"data/clean/{f}", "Description": desc, "Rows": len(df), "Columns": df.shape[1],
                   "Unit of observation": {"national_monthly.csv": "month", "regional_panel.csv": "region x month",
                                           "weights_2018.csv": "basket x code x region x item",
                                           "region_info.csv": "region", "items_2026.csv": "basket x item x month"}[f]})
    for s in SOURCES:
        for path in s["files"]:
            if path.startswith(("psa/", "wb/", "bsp/", "bis/", "fao/", "extracts/")):
                full = RAW / path
                rows = sum(1 for _ in open(full, "rb")) - 1 if full.suffix == ".csv" else "(Excel/PDF)"
                ds.append({"File": f"data/raw/{path}", "Description": f"Raw, {s['id']}: {s['title']}", "Rows": rows,
                           "Columns": "", "Unit of observation": "as published"})
    ds = pd.DataFrame(ds)

    raw_vars = []
    for s in SOURCES:
        for path in s["files"]:
            p = RAW / path
            if p.suffix != ".csv" or not p.exists():
                continue
            cols = pd.read_csv(p, nrows=0, encoding="latin-1").columns
            months = [c for c in cols if re.match(r"^\d{4} [A-Z][a-z]{2}$", c)]
            aves = [c for c in cols if re.match(r"^\d{4} Ave$", c)]
            for c in cols:
                if c in months or c in aves:
                    continue
                raw_vars.append({"File": f"data/raw/{path}", "Source": s["id"], "Column": c,
                                 "Description": "as published by the source"})
            if months:
                raw_vars.append({"File": f"data/raw/{path}", "Source": s["id"],
                                 "Column": f"'{months[0]}' ... '{months[-1]}' ({len(months)} columns)",
                                 "Description": "Monthly index values, 2018 = 100, one column per month"})
            if aves:
                raw_vars.append({"File": f"data/raw/{path}", "Source": s["id"], "Column": f"'{aves[0]}' ... ({len(aves)} columns)",
                                 "Description": "Annual averages (not used)"})
    raw_vars = pd.DataFrame(raw_vars)

    cv = []
    for f, df in data.items():
        for c in df.columns:
            lab, unit, src, tr, rq = describe(c)
            first, last = coverage(df, c)
            cv.append({"Dataset": f, "Variable": c, "Label": lab, "Units": unit, "Source": src,
                       "Transformation": tr, "Used in": rq, "First month": first, "Last month": last,
                       "Non-missing": int(df[c].notna().sum()), "Type": "numeric" if pd.api.types.is_numeric_dtype(df[c]) else "text"})
    cv = pd.DataFrame(cv)

    src = pd.DataFrame([{"ID": s["id"], "Publisher": s["publisher"], "Title": s["title"], "Used for": s["use"],
                         "Files (data/raw/)": "; ".join(s["files"]), "Links": "\n".join(s["links"]),
                         "Coverage": s["coverage"], "Licence / terms": s["licence"], "Downloaded": "2026-10-04"}
                        for s in SOURCES])

    reg = pd.DataFrame([
        {"Code": k, "Name": v, "PSA label variants seen in the raw files": lab, "Treatment": tr}
        for k, v, lab, tr in [
            ("PH", REGION_NAMES["PH"], "PHILIPPINES", "National series"),
            ("NCR", REGION_NAMES["NCR"], "National Capital Region (NCR); NCR", "Unchanged"),
            ("CAR", REGION_NAMES["CAR"], "Cordillera Administrative Region (CAR); CAR", "Unchanged"),
            ("R01", REGION_NAMES["R01"], "Region I (Ilocos Region)", "Unchanged"),
            ("R02", REGION_NAMES["R02"], "Region II (Cagayan Valley)", "Unchanged"),
            ("R03", REGION_NAMES["R03"], "Region III (Central Luzon)", "Unchanged"),
            ("R04A", REGION_NAMES["R04A"], "Region IV-A (CALABARZON)", "Unchanged"),
            ("R04B", REGION_NAMES["R04B"], "MIMAROPA Region; Region IV-B", "Unchanged"),
            ("R05", REGION_NAMES["R05"], "Region V (Bicol Region)", "Unchanged"),
            ("R06", REGION_NAMES["R06"], "Region VI (Western Visayas)", "Old code includes Negros Occidental; "
             "Jan-Aug 2026 chained from new-code changes and flagged new_geo_ext"),
            ("R07", REGION_NAMES["R07"], "Region VII (Central Visayas)", "Old code includes Negros Oriental and "
             "Siquijor; Jan-Aug 2026 chained and flagged"),
            ("R08", REGION_NAMES["R08"], "Region VIII (Eastern Visayas)", "Unchanged"),
            ("R09", REGION_NAMES["R09"], "Region IX (Zamboanga Peninsula)", "Unchanged"),
            ("R10", REGION_NAMES["R10"], "Region X (Northern Mindanao)", "Unchanged"),
            ("R11", REGION_NAMES["R11"], "Region XI (Davao Region)", "Unchanged"),
            ("R12", REGION_NAMES["R12"], "Region XII (SOCCSKSARGEN); Reg12", "Jan-Aug 2026 chained and flagged"),
            ("R13", REGION_NAMES["R13"], "Region XIII (Caraga)", "Unchanged"),
            ("BARMM", REGION_NAMES["BARMM"], "Bangsamoro Autonomous Region in Muslim Mindanao (BARMM); ARMM",
             "One series from PSA; whether pre-2019 values use ARMM or BARMM boundaries (Cotabato City) is "
             "unconfirmed. Jan-Aug 2026 chained and flagged"),
            ("NIR", "Negros Island Region", "Negros Island Region (NIR)", "New code only (from 2024); not used "
             "as a separate unit, because the long history exists only on the old code"),
            ("AONCR", "Areas Outside NCR", "Areas Outside National Capital Region (NCR)", "Aggregate; not used"),
        ]])

    res = pd.DataFrame([{"Output (output/tables or output/figures)": o, "Content": c, "Sources": s,
                         "Python": py, "Stata": st, "Placement in the papers": "main text" if role == "core" else "appendix"}
                        for o, c, s, py, st, role in RESULTS])

    out = ROOT / "codebook" / "codebook.xlsx"
    out.parent.mkdir(exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in [("README", readme), ("Datasets", ds), ("Raw_variables", raw_vars),
                         ("Clean_variables", cv), ("Sources", src), ("Regions", reg), ("Results_map", res)]:
            df.to_excel(xw, sheet_name=name, index=False)
            autosize(xw.sheets[name])
    print("wrote", out, "| clean variables:", len(cv), "| unlabelled:", int((cv.Label == "(see source)").sum()))


if __name__ == "__main__":
    main()
